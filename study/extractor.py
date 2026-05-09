"""Extract plain text from uploaded study materials."""

from __future__ import annotations

import io


SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n\n".join(p for p in parts if p.strip())


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # Also pull cell text from any tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n\n".join(paragraphs)


def _extract_text(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


def extract_text(uploaded_file) -> tuple[str, str | None]:
    """Pull text out of a Streamlit UploadedFile. Returns (text, error_message)."""
    name = (uploaded_file.name or "").lower()
    data = uploaded_file.read()

    try:
        if name.endswith(".pdf"):
            return _extract_pdf(data), None
        if name.endswith(".docx"):
            return _extract_docx(data), None
        if name.endswith(".txt") or name.endswith(".md"):
            return _extract_text(data), None
    except Exception as exc:  # noqa: BLE001
        return "", f"Couldn't read this {name.rsplit('.', 1)[-1].upper()} file: {exc}"

    return "", (
        f"Unsupported file type: '{uploaded_file.name}'. "
        "Please upload a PDF, DOCX, TXT, or MD file."
    )


# ── Batch extraction for chat-attached files ──────────────────────────


# Hard caps to keep us under the model's context window. Per-file cap
# kicks in for unusually long single documents; the total cap protects
# against a user attaching a stack of large PDFs at once.
MAX_CHARS_PER_FILE = 12_000
MAX_CHARS_TOTAL = 30_000


def extract_attachments(files) -> tuple[list[tuple[str, str]], list[str]]:
    """Extract text from a list of Streamlit uploaded files.

    Returns (extracted, errors) where:
      - extracted is a list of (filename, text) tuples
      - errors is a list of human-readable error strings (one per problem file)
    """
    extracted: list[tuple[str, str]] = []
    errors: list[str] = []
    if not files:
        return extracted, errors

    total = 0
    for f in files:
        text, err = extract_text(f)
        if err:
            errors.append(f"{f.name}: {err}")
            continue
        if not text or not text.strip():
            errors.append(
                f"{f.name}: no readable text (scanned image PDF?)"
            )
            continue
        if len(text) > MAX_CHARS_PER_FILE:
            text = (
                text[:MAX_CHARS_PER_FILE]
                + "\n\n[…this file was truncated for length…]"
            )
        if total + len(text) > MAX_CHARS_TOTAL:
            remaining = max(0, MAX_CHARS_TOTAL - total)
            text = (
                text[:remaining]
                + "\n\n[…remaining attachments truncated to fit context…]"
            )
        extracted.append((f.name, text))
        total += len(text)
        if total >= MAX_CHARS_TOTAL:
            break
    return extracted, errors


def build_augmented_message(user_text: str, docs: list[tuple[str, str]]) -> str:
    """Combine the user's typed message with attached document text — the
    string the AI actually sees on the wire."""
    if not docs:
        return user_text
    body = user_text.strip() or "(Please look at the attached document.)"
    pieces = [body, "", "=== ATTACHED DOCUMENT(S) ==="]
    for name, content in docs:
        pieces.append("")
        pieces.append(f"--- {name} ---")
        pieces.append(content)
    return "\n".join(pieces)
